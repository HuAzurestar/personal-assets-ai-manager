"""Currency must have provenance; normalization never invents raw columns."""
import csv
import io
import zipfile
import openpyxl
import xlwt
from pathlib import Path
from zoneinfo import ZoneInfo

import pytest
from backend.parser.statement_parser import parse_statement
from backend.parser.file_import import _read_zip
from fictional_statement_pdf import statement_pdf

FIXTURES = Path(__file__).parent / "fixtures" / "pirc35"


def test_abc_document_currency_does_not_change_raw_row():
    table = list(csv.reader(io.StringIO((FIXTURES / "abc-1.csv").read_text(encoding="utf-8"))))
    header = next(row for row in table if "交易金额" in row)
    index = header.index("币种")
    start = table.index(header)
    for row in table[start:]:
        del row[index]
    output = io.StringIO()
    csv.writer(output).writerows(table)
    doc = parse_statement(output.getvalue().encode(), "mock.csv",
                          source_timezone=ZoneInfo("Asia/Hong_Kong"))
    assert all(row["currency"] == "CNY" for row in doc["rows"])
    assert all("币种" not in row["raw"] for row in doc["rows"])


@pytest.mark.parametrize("document_currency", [False, True])
def test_bank_currency_is_required_or_derived_from_document(document_currency):
    table = list(csv.reader(io.StringIO((FIXTURES / "ccb-2.csv").read_text(encoding="utf-8"))))
    if not document_currency:
        table = [row for row in table if not row[0].startswith("币种：")]
    header = next(row for row in table if "币别" in row)
    currency_column = header.index("币别")
    after_header = False
    for row in table:
        if row is header:
            after_header = True
        elif after_header:
            row[currency_column] = ""
    output = io.StringIO()
    csv.writer(output).writerows(table)
    doc = parse_statement(output.getvalue().encode(), "mock.csv",
                          source_timezone=ZoneInfo("Asia/Hong_Kong"))
    if document_currency:
        assert all(row["currency"] == "CNY" and row["raw"]["币别"] == "" for row in doc["rows"])
    else:
        assert all(row["error"] and row["disposition"] == "error" for row in doc["rows"])


@pytest.mark.parametrize("name", ["../ignored/", "C:/statement.csv", "..\\statement.csv"])
def test_zip_rejects_unsafe_paths_even_for_directory_members(name):
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        archive.writestr(name, b"")
        if name.endswith("/"):
            archive.writestr("statement.csv", b"mock")
    with pytest.raises(ValueError, match="unsupported"):
        _read_zip(output.getvalue(), None)


@pytest.mark.parametrize("provider", ["abc", "cmb"])
def test_fictional_text_pdf_layout_parses_all_pages_with_stable_row_numbers(provider):
    content = statement_pdf(provider)
    doc = parse_statement(content, "fictional.pdf", source_timezone=ZoneInfo("Asia/Hong_Kong"))
    assert doc["source_type"] == provider and doc["format"] == "pdf"
    assert [row["row_number"] for row in doc["rows"]] == [1, 2]
    assert [row["raw"]["_page"] for row in doc["rows"]] == ["1", "2"]
    assert all(row["error"] is None and row["amount_minor"] == -1025 and row["currency"] == "CNY" for row in doc["rows"])


def test_pdf_password_is_forwarded_only_to_reader(monkeypatch):
    import backend.parser.statement_parser as module
    original = module.pdfplumber.open
    observed = []
    def reader(*args, **kwargs):
        observed.append(kwargs.pop("password", None))
        return original(*args, **kwargs)
    monkeypatch.setattr(module.pdfplumber, "open", reader)
    doc = parse_statement(statement_pdf(), "fictional.pdf", "fictional-password", source_timezone=ZoneInfo("Asia/Hong_Kong"))
    assert observed == ["fictional-password"]
    assert "fictional-password" not in str(doc)


@pytest.mark.parametrize("extension", ["xls", "xlsx"])
def test_all_workbook_sheets_keep_linear_locations_and_own_account(extension):
    table = list(csv.reader(io.StringIO((FIXTURES / "ccb-2.csv").read_text(encoding="utf-8"))))
    book = openpyxl.Workbook() if extension == "xlsx" else xlwt.Workbook()
    if extension == "xlsx":
        book.remove(book.active)
    for index in range(2):
        sheet = book.create_sheet(f"Mock{index}") if extension == "xlsx" else book.add_sheet(f"Mock{index}")
        for number, original in enumerate(table):
            values = [value.replace("990000000000001234", f"99000000000000123{index}") for value in original]
            if extension == "xlsx":
                sheet.append(values)
            else:
                for column, value in enumerate(values):
                    sheet.write(number, column, value)
    output = io.BytesIO()
    book.save(output)
    doc = parse_statement(output.getvalue(), f"mock.{extension}", source_timezone=ZoneInfo("Asia/Hong_Kong"))
    assert len(doc["rows"]) == 48 and all(row["error"] is None for row in doc["rows"])
    assert [row["row_number"] for row in doc["rows"]] == list(range(6, 30)) + list(range(35, 59))
    assert {row["account"]["number"] for row in doc["rows"][:24]} == {"990000000000001230"}
    assert {row["account"]["number"] for row in doc["rows"][24:]} == {"990000000000001231"}
